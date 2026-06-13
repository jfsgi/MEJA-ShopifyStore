#!/usr/bin/env python3
"""
build_docx.py — Generate styled Word (.docx) versions of the MEJA plan docs.

Reads the Markdown in ../docs and writes .docx files into ../word:
  - One .docx per source document.
  - One combined "MEJA-WebStore-Complete-Plan.docx" with a cover page.

Supports the Markdown subset used in these docs: ATX headings, paragraphs with
inline **bold** / *italic* / `code` / [links], bullet & numbered lists, pipe
tables, fenced code blocks (```), blockquotes, and horizontal rules. Mermaid
code blocks are rendered as captioned monospaced source (diagrams render on
GitHub; in Word they appear as readable source).

Usage:
    pip3 install python-docx
    python3 tools/build_docx.py
"""

import os
import re
import sys

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DOCS_DIR = os.path.join(ROOT, "docs")
OUT_DIR = os.path.join(ROOT, "word")

# Source docs, in order, with friendly titles for the combined cover/TOC.
SOURCES = [
    ("01-Master-Plan.md", "Master Plan"),
    ("02-UI-Design-Standard.md", "UI Design Standard (incl. 10 iterations)"),
    ("03-Product-Workflows.md", "Product Workflows"),
    ("05-Integration-Context.md", "Integration Context (sibling apps)"),
]

MONO = "Consolas"
INLINE_RE = re.compile(
    r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*]+\*|(?<![A-Za-z0-9])_[^_]+_(?![A-Za-z0-9]))"
)
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


# ---------- low-level styling helpers ----------

def shade(paragraph, fill="F4F4F2"):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def add_hr(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BBBBBB")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_formatted_runs(paragraph, text):
    """Add runs to a paragraph, honoring **bold**, *italic*, `code`, [links]."""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("**") and tok.endswith("**"):
            r = paragraph.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = paragraph.add_run(tok[1:-1]); r.font.name = MONO; r.font.size = Pt(9.5)
        elif tok.startswith("["):
            lm = LINK_RE.match(tok)
            label = lm.group(1) if lm else tok
            r = paragraph.add_run(label)
            r.font.color.rgb = RGBColor(0x1A, 0x4F, 0x8A)
            r.underline = True
        elif tok.startswith("*") and tok.endswith("*"):
            r = paragraph.add_run(tok[1:-1]); r.italic = True
        elif tok.startswith("_") and tok.endswith("_"):
            r = paragraph.add_run(tok[1:-1]); r.italic = True
        else:
            paragraph.add_run(tok)
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


# ---------- block parsing ----------

def is_table_sep(line):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return len(cells) > 0 and all(re.fullmatch(r":?-{2,}:?", c or "") for c in cells)


def split_cells(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc, rows):
    header = split_cells(rows[0])
    body = [split_cells(r) for r in rows[2:]]
    ncols = len(header)
    table = doc.add_table(rows=1, cols=ncols)
    try:
        table.style = "Light Grid Accent 1"
    except Exception:
        try:
            table.style = "Table Grid"
        except Exception:
            pass
    table.autofit = True
    hdr = table.rows[0].cells
    for i, cell_text in enumerate(header):
        if i >= ncols:
            break
        p = hdr[i].paragraphs[0]
        add_formatted_runs(p, cell_text.replace("<br/>", "\n").replace("<br>", "\n"))
        for run in p.runs:
            run.bold = True
    for row in body:
        cells = table.add_row().cells
        for i in range(ncols):
            txt = row[i] if i < len(row) else ""
            p = cells[i].paragraphs[0]
            # support <br/> as in-cell line breaks
            parts = re.split(r"<br/?>", txt)
            for j, part in enumerate(parts):
                if j > 0:
                    p.add_run().add_break()
                add_formatted_runs(p, part)
    doc.add_paragraph()  # spacing after table


def add_code_block(doc, lang, lines):
    if lang:
        cap = doc.add_paragraph()
        r = cap.add_run(
            ("Diagram (Mermaid source — renders as a diagram on GitHub):"
             if lang.lower() == "mermaid"
             else f"Code ({lang}):")
        )
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    p = doc.add_paragraph()
    shade(p)
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.right_indent = Inches(0.15)
    for i, ln in enumerate(lines):
        if i > 0:
            p.add_run().add_break()
        run = p.add_run(ln if ln else " ")
        run.font.name = MONO
        run.font.size = Pt(9)


HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
ULIST_RE = re.compile(r"^(\s*)[-*]\s+(.*)$")
OLIST_RE = re.compile(r"^(\s*)\d+\.\s+(.*)$")
FENCE_RE = re.compile(r"^```(\w*)\s*$")


def render_markdown(doc, md_text):
    lines = md_text.splitlines()
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]

        # fenced code block
        fm = FENCE_RE.match(line.strip())
        if fm:
            lang = fm.group(1)
            block = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1  # skip closing fence
            add_code_block(doc, lang, block)
            continue

        # blank line
        if not line.strip():
            i += 1
            continue

        # horizontal rule
        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", line.strip()):
            add_hr(doc)
            i += 1
            continue

        # heading
        hm = HEADING_RE.match(line)
        if hm:
            level = len(hm.group(1))
            text = hm.group(2).strip()
            style = "Title" if level == 1 else f"Heading {min(level, 4)}"
            try:
                p = doc.add_paragraph(style=style)
            except Exception:
                p = doc.add_paragraph()
            add_formatted_runs(p, text)
            i += 1
            continue

        # table (header line followed by separator line)
        if line.strip().startswith("|") and i + 1 < n and is_table_sep(lines[i + 1]):
            tbl = [lines[i], lines[i + 1]]
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                tbl.append(lines[i])
                i += 1
            add_table(doc, tbl)
            continue

        # blockquote
        if line.lstrip().startswith(">"):
            text = re.sub(r"^\s*>\s?", "", line)
            p = doc.add_paragraph(style="Intense Quote") if _has_style(doc, "Intense Quote") else doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_formatted_runs(p, text)
            for run in p.runs:
                run.italic = True
            i += 1
            continue

        # unordered list
        um = ULIST_RE.match(line)
        if um:
            indent = len(um.group(1))
            lvl = min(indent // 2, 2)
            style = "List Bullet" if lvl == 0 else f"List Bullet {lvl + 1}"
            p = _styled_paragraph(doc, style)
            add_formatted_runs(p, um.group(2))
            i += 1
            continue

        # ordered list
        om = OLIST_RE.match(line)
        if om:
            p = _styled_paragraph(doc, "List Number")
            add_formatted_runs(p, om.group(2))
            i += 1
            continue

        # plain paragraph
        p = doc.add_paragraph()
        add_formatted_runs(p, line.strip())
        i += 1


def _has_style(doc, name):
    try:
        return name in [s.name for s in doc.styles]
    except Exception:
        return False


def _styled_paragraph(doc, style):
    try:
        return doc.add_paragraph(style=style)
    except Exception:
        return doc.add_paragraph()


# ---------- document assembly ----------

def base_document():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for section in doc.sections:
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
    return doc


def add_footer(doc, text):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = text
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)


def cover_page(doc, title, subtitle):
    for _ in range(4):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("MEJA DESIGNS")
    r.bold = True
    r.font.size = Pt(30)
    r.font.color.rgb = RGBColor(0x6B, 0x4A, 0x2B)
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = s.add_run(title)
    rs.font.size = Pt(18)
    rs.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    if subtitle:
        ss = doc.add_paragraph()
        ss.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rss = ss.add_run(subtitle)
        rss.italic = True
        rss.font.size = Pt(11)
        rss.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    for _ in range(2):
        doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rm = meta.add_run("Draft for approval · 2026-06-13 · Branch: claude/awesome-cori-hk0m0c")
    rm.font.size = Pt(9)
    rm.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    doc.add_page_break()


def build_individual():
    outputs = []
    for fname, title in SOURCES:
        src = os.path.join(DOCS_DIR, fname)
        if not os.path.exists(src):
            print(f"  ! skip (missing): {fname}")
            continue
        with open(src, encoding="utf-8") as f:
            md = f.read()
        doc = base_document()
        add_footer(doc, "MEJA Designs WebStore — Draft for approval")
        render_markdown(doc, md)
        out = os.path.join(OUT_DIR, os.path.splitext(fname)[0] + ".docx")
        doc.save(out)
        outputs.append(out)
        print(f"  ✓ {os.path.relpath(out, ROOT)}")
    return outputs


def build_combined():
    doc = base_document()
    add_footer(doc, "MEJA Designs WebStore — Complete Plan (Draft for approval)")
    cover_page(doc, "Shopify WebStore — Complete Plan",
               "Master Plan · UI Design Standard (10 iterations) · Product Workflows")
    # simple contents list
    h = doc.add_paragraph(style="Heading 1") if _has_style(doc, "Heading 1") else doc.add_paragraph()
    h.add_run("Contents")
    for idx, (_, title) in enumerate(SOURCES, 1):
        p = doc.add_paragraph(style="List Number") if _has_style(doc, "List Number") else doc.add_paragraph()
        p.add_run(title)
    doc.add_page_break()

    for idx, (fname, _) in enumerate(SOURCES):
        src = os.path.join(DOCS_DIR, fname)
        if not os.path.exists(src):
            continue
        with open(src, encoding="utf-8") as f:
            md = f.read()
        render_markdown(doc, md)
        if idx < len(SOURCES) - 1:
            doc.add_page_break()
    out = os.path.join(OUT_DIR, "MEJA-WebStore-Complete-Plan.docx")
    doc.save(out)
    print(f"  ✓ {os.path.relpath(out, ROOT)}")
    return out


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    print("Building individual Word documents:")
    build_individual()
    print("Building combined Word document:")
    build_combined()
    print("Done.")


if __name__ == "__main__":
    sys.exit(main())
